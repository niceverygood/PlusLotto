#!/usr/bin/env python3
"""Read-only, reproducible cutover sample workbook; private output only.

GET requests only. No member update, hold release, SMS, or external delivery.
python scripts/make-cutover-review.py --env-file /private/.env.local --out /private/review
Use --reuse only to render the saved snapshot; refresh after operational changes.
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import hashlib
import json
import os
import re
import time
import urllib.error
import urllib.parse
import urllib.request
from collections import Counter, defaultdict
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
PROJECT = 'https://xmfdbmlpvvqqkhqemfay.supabase.co'
SITES = {'lotto815': '815로또', 'infolotto': '인포로또', 'cplotto': '일행로또', 'best': '프리미엄로또'}
GRADES = {'simple': '간편가입', 'free': '무료', 'gold': '미정', 'goldp': '실버', 'vip': '골드', 'royal': '다이아', 'ovr': '인반언스', 'toss': '토스DB'}
PAID = {'gold', 'goldp', 'vip', 'royal'}
DAYS = ['일', '월', '화', '수', '목', '금', '토']
META_KEYS = ['source_site', 'legacy_idx', 'legacy_id', 'legacy_level_num', 'legacy_item_code', 'legacy_sales_idx', 'legacy_member_start_datetime', 'legacy_member_end_datetime', 'legacy_primary_weekday', 'legacy_secondary_weekday', 'legacy_secondary_hour', 'legacy_weekly_reco_hour', 'weekly_reco_day', 'weekly_reco_count', 'end_date', 'reco_paused', 'reco_pause_reason', 'legacy_original_phone', 'legacy_consent_review_required', 'legacy_status']


def write_json(path: Path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding='utf-8')
    path.chmod(0o600)


def load_env(path: Path):
    config = {}
    for line in path.read_text().splitlines():
        line = line.strip()
        if not line or line.startswith('#') or '=' not in line:
            continue
        key, value = line.removeprefix('export ').split('=', 1)
        config[key.strip()] = value.strip().strip('\"\'')
    url = config.get('SUPABASE_URL') or config.get('VITE_SUPABASE_URL')
    if url and url.rstrip('/') != PROJECT:
        raise ValueError('Wrong project; expected PlusLotto project only')
    return config['SUPABASE_SERVICE_ROLE_KEY']


class Reader:
    def __init__(self, key):
        self.key = key

    def get(self, table, params):
        url = PROJECT + '/rest/v1/' + table + '?' + urllib.parse.urlencode(params)
        request = urllib.request.Request(url, method='GET', headers={'apikey': self.key, 'Authorization': 'Bearer ' + self.key})
        for attempt in range(3):
            try:
                with urllib.request.urlopen(request, timeout=90) as response:
                    return json.load(response)
            except (urllib.error.HTTPError, urllib.error.URLError) as error:
                if attempt == 2:
                    raise RuntimeError('Read-only query failed: ' + table + ' ' + str(getattr(error, 'code', 'network'))) from None
                time.sleep(1 + attempt)

    def scan(self, table, select, filters=None):
        rows, cursor = [], None
        while True:
            params = {'select': select, 'order': 'id.asc', 'limit': 1000, **(filters or {})}
            if cursor is not None:
                params['id'] = 'gt.' + cursor
            page = self.get(table, params)
            rows.extend(page)
            if len(page) < 1000:
                break
            cursor = page[-1]['id']
        if len(rows) != len({row['id'] for row in rows}):
            raise ValueError('Duplicate scan keys')
        return rows


def snapshot(key):
    reader = Reader(key)
    started = dt.datetime.now(dt.timezone.utc).isoformat()
    members = read_members(reader)
    phone_keys = reader.scan('members', 'id,phone,site:meta->>source_site')
    payments = reader.scan('payments', 'id,member_id,amount,status,method,paid_at,period_start,period_end,product_id,legacy_idx:meta->legacy_idx', {'meta->>source_site': 'in.(' + ','.join(SITES) + ')'})
    staff = reader.scan('staff', 'id,name,role')
    settings = reader.get('site_settings', {'select': 'weekly_free_reco,sms,membership_tiers', 'id': 'eq.1'})[0]
    # Store only required switches, never credentials or full SMS settings.
    sms = settings.get('sms') or {}
    flags = {'oneshot_enabled': bool(sms.get('oneshot_enabled')), 'global_sender_configured': bool(sms.get('sender_no')), 'paid_sms': bool((settings.get('weekly_free_reco') or {}).get('paid_sms')), 'auto_enabled': (settings.get('weekly_free_reco') or {}).get('enabled') is not False, 'senders': {site: bool(re.sub(r'\D', '', str(((sms.get('by_site') or {}).get(site) or {}).get('sender_no') or ''))) for site in SITES}}
    return {'started_at': started, 'finished_at': dt.datetime.now(dt.timezone.utc).isoformat(), 'project': PROJECT, 'members': members, 'phone_keys': phone_keys, 'payments': payments, 'staff': staff, 'settings': flags, 'membership_tiers': settings.get('membership_tiers')}


def read_members(reader):
    select = 'id,user_id,name,phone,grade,status,assigned_staff_id,is_suspended,is_deleted,is_withdrawn,' + ','.join(key + ':meta->' + key for key in META_KEYS)
    return reader.scan('members', select, {'meta->>source_site': 'in.(' + ','.join(SITES) + ')'})


def number(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def parsed_end(value):
    if not isinstance(value, str):
        return None
    match = re.match(r'^(\d{4})-(\d{2})-(\d{2})(?:$|[T\s])', value.strip())
    if not match:
        return None
    try:
        return dt.date(*map(int, match.groups())).isoformat()
    except ValueError:
        return None


def first_day(row, start):
    day = row.get('weekly_reco_day')
    day = day if number(day) else 5 if row['grade'] == 'free' else None
    if day is None or day != int(day) or not 0 <= day < 7:
        return None, None
    start_day = (start.weekday() + 1) % 7
    return int(day), (start + dt.timedelta(days=(day - start_day) % 7)).isoformat()


def assess(row, start, flags):
    day, date = first_day(row, start)
    reasons = []
    if any(row.get(key) is not False for key in ('is_suspended', 'is_deleted', 'is_withdrawn')):
        reasons.append('정지·삭제·탈퇴 또는 상태 플래그 누락')
    if day is None:
        reasons.append('발송요일 미지정/유효하지 않음')
    import_hold = row.get('reco_paused') is True and row.get('reco_pause_reason') == 'legacy_import_review'
    if row.get('reco_paused') is True and not import_hold:
        reasons.append('운영자 설정 일시정지')
    end = parsed_end(row.get('end_date'))
    if date and end and end < date:
        reasons.append('첫 예정일 전에 이용기간 만료')
    if number(row.get('weekly_reco_count')) and row.get('weekly_reco_count') == 0:
        reasons.append('조합수 0')
    paid_on = flags['paid_sms'] and flags['oneshot_enabled'] and flags['global_sender_configured']
    if not flags['auto_enabled'] and not (paid_on and row['grade'] in PAID):
        reasons.append('자동발급 설정 꺼짐')
    if reasons:
        verdict = '발급·발송 제외'
    elif row['grade'] not in PAID:
        verdict = '조합 발급만'
    elif not row.get('phone'):
        verdict = '문자 제외(연락처 없음)'
        reasons.append('휴대폰 없음')
    else:
        verdict = '회원 조건 충족'
    blocks = []
    if row.get('reco_paused') is True:
        blocks.append('이관 검수 보류' if import_hold else '운영자 일시정지')
    if row['grade'] in PAID and not paid_on:
        blocks.append('공통 문자 설정 미충족')
    if row['grade'] in PAID and not flags['senders'][row['source_site']]:
        blocks.append('사이트 발신번호 미등록')
    issues = []
    if row['grade'] in PAID and not end:
        issues.append('이용 종료일 없음/형식 확인')
    if not row.get('assigned_staff_id'):
        issues.append('담당자 미배정')
    if not re.fullmatch(r'01\d{8,9}', re.sub(r'\D', '', row.get('phone') or '')):
        issues.append('연락처 형식 확인')
    if row.get('legacy_secondary_weekday') not in (None, '', '0', 0, '-', 'none'):
        issues.append('원본 추가요일 확인(현재 크론은 주요일만 처리)')
    if row.get('legacy_consent_review_required') is True:
        issues.append('원본 수신동의 정보 확인')
    if day is None:
        issues.append('발송요일 확인')
    if row.get('cross_sites'):
        issues.append('동일번호 다른 사이트 계약 별도 검수')
    return {'day': day, 'first_date': date, 'verdict': verdict, 'reason': '; '.join(reasons), 'current_blocks': '; '.join(blocks), 'issues': '; '.join(issues)}


def select_sample(rows, size=30):
    """Coverage sample, not a population-rate estimator; deterministic salted hash tie break."""
    tokens = {}
    for row in rows:
        a = row['assessment']
        tokens[row['id']] = {('grade', row['grade']), ('day', str(a['day'])), ('verdict', a['verdict']), ('state', row['status']), ('cross', bool(row['cross_sites']))}
        if row.get('legacy_secondary_weekday') not in (None, '', '0', 0, '-', 'none'):
            tokens[row['id']].add(('extra_day', row['id']))
        if row.get('assigned_staff_id'):
            tokens[row['id']].add(('assigned', True))
        if row['source_site'] == 'lotto815' and row['grade'] == 'goldp':
            tokens[row['id']].add(('815_silver', True))
    frequency = Counter(token for values in tokens.values() for token in values)
    ordered = sorted(rows, key=lambda row: hashlib.sha256(('cutover-20261006|' + row['id']).encode()).hexdigest())
    chosen, covered = [], set()
    remaining = ordered.copy()
    while remaining and len(chosen) < size:
        def score(row):
            unseen = tokens[row['id']] - covered
            return sum(1 + 1 / frequency[token] for token in unseen)
        best = max(remaining, key=score)
        if score(best) == 0:
            # Fill by grade/day strata round-robin, so a large grade does not consume every row.
            used = Counter((r['grade'], r['assessment']['day']) for r in chosen)
            best = min(remaining, key=lambda r: used[(r['grade'], r['assessment']['day'])])
        chosen.append(best)
        remaining.remove(best)
        covered |= tokens[best['id']]
    return sorted(chosen, key=lambda row: (row['grade'], str(row['assessment']['day']), row['id']))


def safe_text(value):
    if value is None:
        return ''
    if isinstance(value, (dict, list)):
        value = json.dumps(value, ensure_ascii=False)
    if isinstance(value, str) and value.startswith(('=', '+', '-', '@')):
        return "'" + value
    return value


def render(data, out, start):
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.utils import get_column_letter

    phones = defaultdict(set)
    for key in data['phone_keys']:
        phone = re.sub(r'\D', '', key.get('phone') or '')
        if phone:
            phones[phone].add(key.get('site') or 'pluslotto')
    payments = defaultdict(list)
    for payment in data['payments']:
        payments[payment['member_id']].append(payment)
    staff = {row['id']: row['name'] for row in data['staff']}
    labels = dict(GRADES)
    tiers = data.get('membership_tiers')
    if isinstance(tiers, list):
        labels.update({tier['grade']: tier['label'] for tier in tiers if isinstance(tier, dict) and tier.get('grade') and tier.get('label')})
    for row in data['members']:
        row['cross_sites'] = sorted(phones[re.sub(r'\D', '', row.get('phone') or '')] - {row['source_site']})
        row['assessment'] = assess(row, start, data['settings'])

    wb = Workbook()
    cover = wb.active
    cover.title = '안내'
    observed = dt.datetime.fromisoformat(data['finished_at']).astimezone(dt.timezone(dt.timedelta(hours=9))).strftime('%Y-%m-%d %H:%M KST')
    notes = [
        ['전산 통합 전환 검수표', '4개 사이트 각 30명 / 총 120명'],
        ['운영 DB 조회 완료', observed],
        ['전환 기준', start.isoformat() + ' 이후 첫 지정요일(10/6~10/12)'],
        ['현장 검수', '회원·계약·요일·조합수·종료일·결제·담당을 구전산과 대조 후 결과/수정요청 입력'],
        ['판정 의미', '이관 검수 보류만 해제했다고 가정한 회원 조건. 실제 발송 승인이나 발송 성공 보장이 아님'],
        ['현재 차단', '발신번호·문자 설정·보류를 별도 표기. 보류 해제/설정 변경/문자 발송은 수행하지 않음'],
        ['범위', '운영 DB에 이미 들어간 회원 표본이며, 미이관 격리 자료는 포함하지 않음'],
        ['표본 방법', '등급·요일·상태·기간·동일번호 타사이트 계약·추가요일을 포함한 결정적 커버리지 표본'],
        ['해석 주의', '표본은 전체 오류율 추정용 무작위 표본이 아님. 30명 검수로 전체 회원 정상 여부를 보장하지 않음'],
        ['동일번호', '다른 사이트 계약은 서로 별도 회원으로 유지. 동일번호라는 이유로 합치거나 누락시키지 않음'],
        ['크론 기준', 'api/weekly-reco.ts 현재 코드의 정지·기간·요일·조합수·등급·설정 규칙. 실행 당일 이미 발급된 회차는 별도 점검'],
        ['추가요일', '원본 보조요일은 별도 표시. 현재 크론은 weekly_reco_day 한 개만 처리하므로 추가요일 계약을 별도 확인'],
        ['원본 날짜', '원본 날짜/코드는 이관 시 보존된 meta 값이며, 원본 SQL 전체를 이번 조회에서 재파싱하지 않음'],
        ['조회 일관성', '회원/결제는 키셋 페이지별 읽기. 단일 DB 트랜잭션 스냅샷은 아니며 관측 시간 이후 현장 변경이 반영되지 않음'],
        ['공유', '개인정보 포함. 검수 담당자에게만 전달. 이 파일을 Git에 추가하지 않음'],
    ]
    for note in notes:
        cover.append(note)
    cover.sheet_properties.pageSetUpPr.fitToPage = True
    cover.page_setup.orientation = 'landscape'
    cover.page_setup.paperSize = cover.PAPERSIZE_A3
    cover.page_setup.fitToWidth = 1
    cover.page_setup.fitToHeight = 1
    cover.column_dimensions['A'].width = 25
    cover.column_dimensions['B'].width = 115
    for row in cover:
        for cell in row:
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            cell.font = Font(name='Pretendard', size=11)
        cover.row_dimensions[row[0].row].height = 36 if row[0].row > 1 else 30
    cover['A1'].fill = cover['B1'].fill = PatternFill('solid', fgColor='182A52')
    cover['A1'].font = cover['B1'].font = Font(name='Pretendard', color='FFFFFF', bold=True, size=14)

    main_headers = ['번호', '검수결과', '수정요청', '이름', '전화번호', '구전산 ID', '등급', '요일', '첫 예정일', '종료일', '조합수', '회원조건 판정', '제외 사유', '현재 차단', '점검 항목', '동일번호 타사이트', '담당자', '승인결제 건수', '최종 승인결제일', '회원 키']
    detail_headers = ['번호', '회원 키', '원본 회원번호', '현재 ID', '원본 등급코드', '원본 상품코드', '원본 담당코드', '원본 상태', '현재 상태', '원본 시작일', '원본 종료일', '현재 종료일', '원본 주요일', '원본 추가요일', '원본 시각', '현재 요일값', '현재 조합수', '승인결제 건수', '승인결제 합계(원)', '최종 승인결제일', '최종 결제수단', '최종 결제 시작일', '최종 결제 종료일', '이관 보류', '발신번호 등록']
    summary = []
    chosen_all = []

    def format_sheet(sheet, widths):
        sheet.freeze_panes = 'D2'
        sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_view.zoomScale = 85
        for index, width in enumerate(widths, 1):
            sheet.column_dimensions[get_column_letter(index)].width = width
        for cell in sheet[1]:
            cell.fill = PatternFill('solid', fgColor='182A52')
            cell.font = Font(name='Pretendard', size=11, color='FFFFFF', bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical='center')
        sheet.row_dimensions[1].height = 36
        for row in sheet.iter_rows(min_row=2):
            sheet.row_dimensions[row[0].row].height = 56
            for cell in row:
                cell.font = Font(name='Pretendard', size=10)
                cell.alignment = Alignment(wrap_text=True, vertical='center', horizontal='left', indent=1)
                cell.fill = PatternFill('solid', fgColor='F7F8FA' if cell.row % 2 == 0 else 'FFFFFF')
                cell.border = Border(bottom=Side(style='hair', color='E1E6EC'))
        sheet.print_options.horizontalCentered = True
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = 'landscape'
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.print_title_rows = '1:1'

    for site, label in SITES.items():
        population = [row for row in data['members'] if row['source_site'] == site]
        chosen = select_sample(population)
        if len(chosen) != 30:
            raise ValueError('Expected exactly 30 samples: ' + site)
        chosen_all.extend(chosen)
        sheet, detail = wb.create_sheet(label), wb.create_sheet(label + '_상세')
        sheet.append(main_headers)
        detail.append(detail_headers)
        csv_rows = []
        for index, row in enumerate(chosen, 1):
            a = row['assessment']
            approved = sorted([payment for payment in payments[row['id']] if payment['status'] == 'approved'], key=lambda p: p.get('paid_at') or '', reverse=True)
            latest = approved[0] if approved else {}
            main = [index, '미검수', '', row['name'], row.get('phone'), row.get('legacy_id'), labels.get(row['grade'], row['grade']), DAYS[a['day']] if a['day'] is not None else '미지정', a['first_date'], row.get('end_date'), row.get('weekly_reco_count'), a['verdict'], a['reason'], a['current_blocks'], a['issues'], ', '.join(SITES.get(s, '플러스로또' if s == 'pluslotto' else s) for s in row['cross_sites']), staff.get(row.get('assigned_staff_id'), '미배정'), len(approved), latest.get('paid_at'), row['id']]
            details = [index, row['id'], row.get('legacy_idx'), row.get('user_id'), row.get('legacy_level_num'), row.get('legacy_item_code'), row.get('legacy_sales_idx'), row.get('legacy_status'), row.get('status'), row.get('legacy_member_start_datetime'), row.get('legacy_member_end_datetime'), row.get('end_date'), row.get('legacy_primary_weekday'), row.get('legacy_secondary_weekday'), row.get('legacy_weekly_reco_hour'), row.get('weekly_reco_day'), row.get('weekly_reco_count'), len(approved), sum(payment['amount'] for payment in approved), latest.get('paid_at'), latest.get('method'), latest.get('period_start'), latest.get('period_end'), row.get('reco_paused'), data['settings']['senders'][site]]
            sheet.append([safe_text(value) for value in main])
            detail.append([safe_text(value) for value in details])
            csv_rows.append(dict(zip(main_headers, main)) | dict(zip(detail_headers, details)))
        format_sheet(sheet, [7, 13, 30, 14, 18, 20, 11, 9, 14, 14, 10, 21, 28, 29, 52, 25, 18, 12, 26, 39])
        format_sheet(detail, [7, 39, 15, 20, 14, 20, 14, 14, 14, 24, 24, 14, 14, 16, 14, 12, 12, 13, 21, 26, 16, 26, 26, 12, 16])
        validation = DataValidation(type='list', formula1='"미검수,정상,수정필요,확인보류"')
        sheet.add_data_validation(validation)
        validation.add('B2:B31')
        for cells in sheet.iter_rows(min_row=2, min_col=2, max_col=3):
            for cell in cells:
                cell.fill = PatternFill('solid', fgColor='FFF7E6')
        for cell in sheet['E'][1:]:
            cell.number_format = '@'
        csv_path = out / (label + '_검수표_30명.csv')
        with csv_path.open('w', newline='', encoding='utf-8-sig') as handle:
            writer = csv.DictWriter(handle, fieldnames=list(csv_rows[0]))
            writer.writeheader()
            writer.writerows({key: safe_text(value) for key, value in row.items()} for row in csv_rows)
        summary.append({'site': site, 'label': label, 'population': len(population), 'sample': len(chosen), 'population_verdicts': dict(Counter(row['assessment']['verdict'] for row in population)), 'sample_verdicts': dict(Counter(row['assessment']['verdict'] for row in chosen)), 'population_grades': dict(Counter(row['grade'] for row in population)), 'sample_grades': dict(Counter(row['grade'] for row in chosen)), 'population_days': dict(Counter(str(row['assessment']['day']) for row in population)), 'sample_days': dict(Counter(str(row['assessment']['day']) for row in chosen)), 'sample_cross_site': sum(bool(row['cross_sites']) for row in chosen), 'sender_configured': data['settings']['senders'][site]})
    overview = wb.create_sheet('전체요약', 1)
    overview.append(['사이트', 'DB 회원 수', '표본 수', '회원조건 충족', '발급·발송 제외', '조합 발급만', '표본 타사이트 중복계약', '사이트 발신번호', '현재 상태'])
    for item in summary:
        verdicts = item['population_verdicts']
        overview.append([item['label'], item['population'], item['sample'], verdicts.get('회원 조건 충족', 0), verdicts.get('발급·발송 제외', 0), verdicts.get('조합 발급만', 0), item['sample_cross_site'], '등록' if item['sender_configured'] else '미등록', '이관 검수 보류 유지'])
    overview.append(['집계 설명', '운영 DB 스냅샷 관측값', '', '보류 해제만 가정한 회원 조건', '첫 지정일 기준', '', '한 사람의 다른 사이트 계약은 별도', '', '실제 발송 가능 인원/성공 건수 아님'])
    format_sheet(overview, [20, 21, 12, 25, 23, 18, 30, 22, 34])
    workbook = out / '4개사이트_전환검수표_각30명.xlsx'
    wb.save(workbook)
    check = load_workbook(workbook)
    errors = []
    for site, label in SITES.items():
        if check[label].max_row != 31 or check[label + '_상세'].max_row != 31:
            errors.append(label + ' row count mismatch')
        for cell in check[label]['E'][1:]:
            if cell.data_type != 's':
                errors.append(label + ' phone is not text')
        if not any((check[label].cell(index, 18).value or 0) > 0 for index in range(2, 32)):
            errors.append(label + ' sample payment aggregates unexpectedly all zero')
    for sheet in check:
        for row in sheet:
            for cell in row:
                if cell.data_type in ('e', 'f'):
                    errors.append('Unexpected Excel formula/error')
    if len({row['id'] for row in chosen_all}) != 120:
        errors.append('Duplicate sample IDs')
    write_json(out / 'selected-members.private.json', chosen_all)
    qa = {'verified_at': dt.datetime.now(dt.timezone.utc).isoformat(), 'observed': observed, 'cutover_start': start.isoformat(), 'sheets': check.sheetnames, 'samples': 120, 'errors': errors, 'settings': data['settings'], 'sites': summary, 'cron_sha256': hashlib.sha256((ROOT / 'api/weekly-reco.ts').read_bytes()).hexdigest(), 'files': {path.name: hashlib.sha256(path.read_bytes()).hexdigest() for path in [workbook, *out.glob('*_검수표_30명.csv')]}}
    write_json(out / 'quality-report.json', qa)
    (out / '검수_안내.txt').write_text('\n\n'.join(key + ': ' + value for key, value in notes) + '\n', encoding='utf-8')
    print(json.dumps({'workbook': str(workbook), 'samples': 120, 'sites': summary, 'qa_errors': errors}, ensure_ascii=False))
    if errors:
        raise ValueError('Workbook QA failed')


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--env-file', type=Path)
    parser.add_argument('--out', type=Path, required=True)
    parser.add_argument('--reuse', action='store_true')
    parser.add_argument('--refresh-members', action='store_true', help='Reuse non-member reads, then refresh member state after a controlled change')
    parser.add_argument('--start', default='2026-10-06')
    args = parser.parse_args()
    if args.out.resolve().is_relative_to(ROOT):
        raise ValueError('PII output must be outside the repository')
    os.umask(0o077)
    args.out.mkdir(parents=True, exist_ok=True)
    args.out.chmod(0o700)
    path = args.out / 'snapshot.private.json'
    if args.reuse or args.refresh_members:
        data = json.loads(path.read_text())
        if args.refresh_members:
            if not args.env_file:
                parser.error('--env-file required for member refresh')
            data['member_refresh_started_at'] = dt.datetime.now(dt.timezone.utc).isoformat()
            data['members'] = read_members(Reader(load_env(args.env_file)))
            data['finished_at'] = dt.datetime.now(dt.timezone.utc).isoformat()
            write_json(path, data)
    else:
        if not args.env_file:
            parser.error('--env-file required for read-only refresh')
        data = snapshot(load_env(args.env_file))
        write_json(path, data)
    render(data, args.out, dt.date.fromisoformat(args.start))


if __name__ == '__main__':
    main()
